from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
src = root / 'docs/reports/20261021/P0阶段总结与下一阶段安排.md'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(.9)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.3)

def font_style(style, size, font='宋体', bold=False):
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rf = style.element.get_or_add_rPr()
    fonts = rf.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rf.insert(0, fonts)
    for key in ('ascii','hAnsi','eastAsia','cs'):
        fonts.set(qn('w:'+key), font)
    style.paragraph_format.widow_control = True

font_style(doc.styles['Normal'], 11.5)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.25
normal.space_after = Pt(7)
font_style(doc.styles['Title'], 23, '微软雅黑', True)
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.styles['Title'].paragraph_format.space_after = Pt(14)
for name, size in [('Heading 1', 15), ('Heading 2', 12.5)]:
    font_style(doc.styles[name], size, '微软雅黑', True)
    fmt = doc.styles[name].paragraph_format
    fmt.space_before = Pt(12); fmt.space_after = Pt(7)
    fmt.keep_with_next = True
for st in ('List Bullet','List Number'):
    font_style(doc.styles[st], 11.5)
    doc.styles[st].paragraph_format.space_after = Pt(5)
code_style = doc.styles.add_style('Report Code', 1)
font_style(code_style, 9.5, 'Consolas')
code_style.paragraph_format.line_spacing = 1.05
code_style.paragraph_format.space_after = Pt(0)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('P0阶段总结  ·  '); r.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def add_table(rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if len(rows[0]) == 6:
        widths = [1.05, 1.1, 1.1, 1.2, 1.25, 1.2]
    else:
        widths = [1.0,1.9,2.0,2.0] if 'E0' in rows[0][1] else [1.0,1.5,2.1,2.3]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    pr.append(borders)
    for i, rowvals in enumerate(rows):
        row = table.rows[0] if i == 0 else table.add_row()
        trPr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trPr.append(no_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); trPr.append(repeat)
        for j, (cell, value) in enumerate(zip(row.cells, rowvals)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, val in [('top',100),('bottom',100),('left',110),('right',110)]:
                elem=OxmlElement('w:'+side); elem.set(qn('w:w'),str(val)); elem.set(qn('w:type'),'dxa'); margins.append(elem)
            tcPr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '203C56' if i == 0 else ('F3F6F8' if i % 2 else 'FFFFFF')); tcPr.append(shade)
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            r = p.add_run(value); r.font.size = Pt(10.5)
            if i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines = src.read_text(encoding='utf-8').splitlines()
i = 0
in_code = False
while i < len(lines):
    line = lines[i].strip(); i += 1
    if line.startswith('```'):
        in_code = not in_code
        if not in_code:
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        continue
    if in_code:
        p = doc.add_paragraph(lines[i-1], 'Report Code')
        p.paragraph_format.keep_with_next = True
        continue
    if not line: continue
    if line == '<!-- pagebreak -->':
        doc.add_page_break(); continue
    if line.startswith('<!--'):
        continue
    if line.startswith('|'):
        rows = [line]
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append(lines[i].strip()); i += 1
        cells = [[c.strip() for c in row.strip('|').split('|')] for row in rows]
        add_table([cells[0], *cells[2:]])
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title'); continue
    if line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1'); continue
    if line.startswith('- '):
        doc.add_paragraph(line[2:], 'List Bullet'); continue
    m = re.match(r'^(\d+)\. (.*)', line)
    if m:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.18)
        p.paragraph_format.first_line_indent = Inches(-.18)
        p.paragraph_format.space_after = Pt(5)
        p.add_run(m.group(1)+'. '+m.group(2)); continue
    p = doc.add_paragraph(line)
    if line.startswith(('整理日期：','项目：')):
        for r in p.runs: r.font.size = Pt(10)
    if any(line.startswith(v) for v in ('汇总报告：','汇总数据：','原始记录：')):
        for r in p.runs: r.font.size=Pt(9)

doc.core_properties.title = 'P0阶段总结与下一阶段安排'
doc.core_properties.subject = 'P0工作和实验结果汇总及G0与P1安排'
doc.core_properties.language = 'zh-CN'
out = src.with_suffix('.docx')
doc.save(out)
print(out)
